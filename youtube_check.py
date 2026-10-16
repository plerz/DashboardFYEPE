# DashboardFYEPE YouTube link checker
import csv, io, os, re, uuid, json, email.utils
from datetime import datetime, timezone, timedelta
from urllib.parse import urlparse, parse_qs
import requests
from flask import request, render_template, redirect, url_for, flash, abort, Response

# YouTube Data API v3: satu permintaan dapat memeriksa maksimal 50 video.
CHECK_BATCH = 50
API_URL = "https://www.googleapis.com/youtube/v3/videos"

def check_api_videos(video_ids):
    """Return (video_id -> (status, note), error, pause_minutes).

    ID yang tidak dikembalikan API tidak otomatis berarti dihapus; video
    dapat privat, dibatasi, atau tidak tersedia untuk pemanggil API.
    """
    key = os.environ.get("YOUTUBE_API_KEY", "").strip()
    if not key:
        return {}, "YOUTUBE_API_KEY belum diatur di Environment Render.", 0
    try:
        response = requests.get(API_URL, params={
            "part": "id,snippet,status", "id": ",".join(video_ids), "key": key,
            "maxResults": 50}, timeout=(5, 18))
    except requests.RequestException as exc:
        return {}, "Koneksi ke YouTube Data API gagal ("+type(exc).__name__+").", 5
    if response.status_code != 200:
        reason = ""
        try:
            err = response.json().get("error", {})
            reason = str(err.get("message") or "")[:180]
            details = err.get("errors") or []
            if details:
                reason = str(details[0].get("reason") or "")+": "+reason
        except (ValueError, AttributeError, TypeError):
            reason = response.text[:120]
        if response.status_code in (401, 403):
            minutes = 60 if "quota" in reason.lower() else 30
        elif response.status_code == 429:
            minutes = 30
        elif response.status_code >= 500:
            minutes = 10
        else:
            minutes = 0
        return {}, "YouTube Data API HTTP %s — %s" % (response.status_code, reason or "Permintaan ditolak"), minutes
    try:
        data = response.json()
        found = {item["id"]: item for item in data.get("items", []) if item.get("id")}
    except (ValueError, TypeError, KeyError):
        return {}, "Respons YouTube Data API tidak dapat dibaca.", 5
    result = {}
    for vid in video_ids:
        item = found.get(vid)
        if not item:
            result[vid] = ("TIDAK DITEMUKAN API", "YouTube Data API tidak mengembalikan video ini. Penyebabnya belum dapat dipastikan (bisa privat, dihapus, atau dibatasi); bukan bukti takedown.")
            continue
        meta = item.get("status") or {}
        privacy = meta.get("privacyStatus", "")
        upload_status = meta.get("uploadStatus", "")
        rejection = meta.get("rejectionReason", "")
        # Hanya penolakan eksplisit dengan alasan penegakan aturan yang dapat
        # diklasifikasikan sebagai takedown terkonfirmasi; bukan hasil kosong.
        if upload_status == "rejected" and rejection in ("copyright", "trademark", "termsOfUse"):
            result[vid] = ("TAKEDOWN TERKONFIRMASI", "API melaporkan uploadStatus=rejected; rejectionReason=" + rejection)
        elif upload_status in ("deleted", "rejected", "failed"):
            result[vid] = ("PERLU VERIFIKASI", "API melaporkan uploadStatus=" + upload_status + ("; rejectionReason=" + rejection if rejection else "") + ". Tidak otomatis berarti takedown oleh YouTube.")
        elif privacy == "public":
            result[vid] = ("TAYANG", "Metadata video publik ditemukan di YouTube Data API; bukan pengujian pemutaran.")
        elif privacy == "unlisted":
            result[vid] = ("TAYANG", "Metadata video unlisted ditemukan; dapat diakses dengan tautan, tetapi tidak ditampilkan dalam pencarian publik.")
        elif privacy == "private":
            result[vid] = ("TIDAK PUBLIK", "YouTube Data API melaporkan privacyStatus=private.")
        else:
            result[vid] = ("PERLU VERIFIKASI", "Video ditemukan di API tetapi status privasinya tidak dapat dipastikan.")
    return result, "", 0

def video_id(raw):
    raw = raw.strip()
    if not raw: return None
    try:
        u = urlparse(raw)
        host = (u.hostname or '').lower()
        if host in ('youtu.be','www.youtu.be'):
            vid = u.path.strip('/').split('/')[0]
        elif host in ('youtube.com','www.youtube.com','m.youtube.com','music.youtube.com'):
            parts = u.path.strip('/').split('/')
            vid = parse_qs(u.query).get('v',[''])[0] if u.path == '/watch' else (parts[1] if len(parts)>1 and parts[0] in ('shorts','live','embed') else '')
        else: return None
        return vid if re.fullmatch(r'[A-Za-z0-9_-]{11}',vid) else None
    except Exception: return None

def parse_txt(content):
    text = content.decode('utf-8-sig', errors='replace')
    output=[]; seen=set()
    for line in text.splitlines():
        if not line.strip(): continue
        cells=line.split('\t')
        match = next(((c.strip(),video_id(c.strip())) for c in reversed(cells) if video_id(c.strip())), None)
        if not match:
            # Allow a line containing just the link, or comma-separated data.
            links = re.findall(r'https?://[^\s,;]+',line)
            match = next(((x,video_id(x)) for x in links if video_id(x)), None)
        if not match: continue
        link, vid=match
        if vid in seen: continue
        seen.add(vid)
        # Preserve source metadata when tabular input is supplied.
        number = cells[0].strip() if len(cells)>1 else str(len(output)+1)
        title = cells[1].strip() if len(cells)>2 else ''
        artist = cells[2].strip() if len(cells)>3 else ''
        source_status = cells[3].strip() if len(cells)>4 else ''
        output.append((number,title,artist,source_status,link,vid))
    return output


def register_youtube_check(app, db, role_required, current_user):
    from flask import jsonify
    def ensure_tables(c):
        c.execute("""CREATE TABLE IF NOT EXISTS youtube_check_rows (
            batch_id TEXT NOT NULL, owner_id TEXT NOT NULL, row_no INTEGER NOT NULL,
            source_no TEXT, song_title TEXT, artist TEXT, source_status TEXT,
            video_url TEXT, video_id TEXT, youtube_status TEXT, checked_at TEXT,
            note TEXT, PRIMARY KEY(batch_id,row_no))""")
        c.execute("""CREATE TABLE IF NOT EXISTS youtube_check_cooldown (
            batch_id TEXT PRIMARY KEY, retry_at TEXT NOT NULL, http_reason TEXT NOT NULL,
            updated_at TEXT NOT NULL)""")
        c.commit()

    def scoped(batch, user):
        if user["role"]=="superadmin":
            return "batch_id=?", (batch,)
        return "batch_id=? AND owner_id=?", (batch,str(user["id"]))

    @app.route("/check-video-youtube",methods=["GET","POST"])
    @role_required("superadmin","admin_pdn")
    def youtube_check_page():
        user=current_user()
        if request.method=="POST":
            upload=request.files.get("video_txt")
            if not upload or not upload.filename.lower().endswith(".txt"):
                flash("Pilih file TXT.","error")
                return redirect(url_for("youtube_check_page"))
            # No fixed number-of-links limit; 32 MB per upload protects server memory.
            content=upload.stream.read(32*1024*1024+1)
            if len(content)>32*1024*1024:
                flash("Ukuran file maksimal 32 MB per upload.","error")
                return redirect(url_for("youtube_check_page"))
            records=parse_txt(content)
            if not records:
                flash("Tidak ada URL YouTube valid di TXT.","error")
                return redirect(url_for("youtube_check_page"))
            batch=uuid.uuid4().hex
            c=db(); ensure_tables(c)
            for i,(number,title,artist,source_status,link,vid) in enumerate(records,1):
                c.execute("""INSERT INTO youtube_check_rows
                    (batch_id,owner_id,row_no,source_no,song_title,artist,source_status,video_url,video_id,youtube_status,checked_at,note)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (batch,str(user["id"]),i,number,title,artist,source_status,link,vid,"BELUM DIPERIKSA","",""))
            c.commit(); c.close()
            return redirect(url_for("youtube_check_page",batch=batch))
        batch=request.args.get("batch","")
        try: page=max(1,int(request.args.get("page","1")))
        except ValueError: page=1
        rows=[]; counts={}; total=0; cooldown=None
        tab=request.args.get("tab","semua")
        if tab not in ("semua","tidak-tersedia","tayang","verifikasi","tidak-publik","tidak-ditemukan","takedown"):
            tab="semua"
        if batch:
            c=db(); ensure_tables(c); where,params=scoped(batch,user)
            # Migrasi aman hasil lama: klasifikasi berbasis halaman watch direset
            # agar dapat diperiksa ulang lewat API. Status awal TXT tidak berubah.
            # Parameterisasi pola LIKE diperlukan untuk PostgreSQL/psycopg2.
            # Karakter % literal dalam SQL yang diberi parameter dapat menyebabkan
            # IndexError: tuple index out of range pada adapter DB.
            legacy_patterns = (
                "Pemutar YouTube%", "Status OK%",
                "YouTube tidak menyediakan%", "Status pemutar%",
            )
            c.execute("""UPDATE youtube_check_rows SET youtube_status='BELUM DIPERIKSA',
                checked_at='',note='Menunggu verifikasi ulang dengan YouTube Data API v3'
                WHERE """ + where + " AND (note LIKE ? OR note LIKE ? OR note LIKE ? OR note LIKE ?)",
                params + legacy_patterns)
            c.commit()
            counts={r["youtube_status"]:r["n"] for r in c.execute(
                "SELECT youtube_status,COUNT(*) n FROM youtube_check_rows WHERE "+where+" GROUP BY youtube_status",params).fetchall()}
            total=sum(counts.values())
            cooldown_row=c.execute("SELECT retry_at,http_reason FROM youtube_check_cooldown WHERE batch_id=?",(batch,)).fetchone()
            if cooldown_row:
                retry_at=datetime.fromisoformat(cooldown_row["retry_at"])
                if retry_at > datetime.now(timezone.utc):
                    cooldown={"retry_at":cooldown_row["retry_at"],"reason":cooldown_row["http_reason"]}
            filter_sql={"tidak-tersedia":" AND youtube_status='TIDAK TERSEDIA'",
                        "tayang":" AND youtube_status='TAYANG'",
                        "verifikasi":" AND youtube_status='PERLU VERIFIKASI'",
                        "tidak-publik":" AND youtube_status='TIDAK PUBLIK'",
                        "tidak-ditemukan":" AND youtube_status='TIDAK DITEMUKAN API'",
                        "takedown":" AND youtube_status='TAKEDOWN TERKONFIRMASI'"}.get(tab, "")
            rows=[dict(r) for r in c.execute(
                "SELECT * FROM youtube_check_rows WHERE "+where+filter_sql+" ORDER BY row_no LIMIT 100 OFFSET ?",
                params+((page-1)*100,)).fetchall()]
            c.close()
            if not total: abort(404)
        return render_template("check_video_youtube.html",user=user,rows=rows,batch=batch,
            counts=counts,total=total,page=page,tab=tab,
            tab_total=({"tidak-tersedia":counts.get("TIDAK TERSEDIA",0),
                        "tayang":counts.get("TAYANG",0),
                        "verifikasi":counts.get("PERLU VERIFIKASI",0),
                        "tidak-publik":counts.get("TIDAK PUBLIK",0),
                        "tidak-ditemukan":counts.get("TIDAK DITEMUKAN API",0),
                        "takedown":counts.get("TAKEDOWN TERKONFIRMASI",0)}.get(tab,total)),
            api_ready=bool(os.environ.get("YOUTUBE_API_KEY", "").strip()),cooldown=cooldown)

    @app.post("/check-video-youtube/reset/<batch>")
    @role_required("superadmin","admin_pdn")
    def youtube_check_reset(batch):
        user=current_user()
        c=db(); ensure_tables(c); where,params=scoped(batch,user)
        found=c.execute("SELECT COUNT(*) n FROM youtube_check_rows WHERE "+where,params).fetchone()["n"]
        if not found: c.close(); abort(404)
        c.execute("UPDATE youtube_check_rows SET youtube_status='BELUM DIPERIKSA',checked_at='',note='Menunggu pemeriksaan melalui YouTube Data API v3' WHERE "+where,params)
        c.execute("DELETE FROM youtube_check_cooldown WHERE batch_id=?",(batch,))
        c.commit(); c.close()
        flash("Status lama direset. Klik Mulai / Lanjutkan untuk memeriksa ulang melalui YouTube Data API v3.","info")
        return redirect(url_for("youtube_check_page",batch=batch))

    @app.post("/check-video-youtube/step/<batch>")
    @role_required("superadmin","admin_pdn")
    def youtube_check_step(batch):
        user=current_user()
        c=db(); ensure_tables(c); where,params=scoped(batch,user)
        total=c.execute("SELECT COUNT(*) n FROM youtube_check_rows WHERE "+where,params).fetchone()["n"]
        if not total: c.close(); abort(404)
        now_dt=datetime.now(timezone.utc)
        cooldown_row=c.execute("SELECT retry_at,http_reason FROM youtube_check_cooldown WHERE batch_id=?",(batch,)).fetchone()
        if cooldown_row and datetime.fromisoformat(cooldown_row["retry_at"]) > now_dt:
            remaining=c.execute("SELECT COUNT(*) n FROM youtube_check_rows WHERE "+where+
                " AND youtube_status='BELUM DIPERIKSA'",params).fetchone()["n"]
            c.close()
            return jsonify(ok=True,done=False,processed=total-remaining,total=total,
                throttled=True,retry_at=cooldown_row["retry_at"],
                warning="Pemeriksaan dijeda otomatis: "+cooldown_row["http_reason"])
        chunk=[dict(r) for r in c.execute("SELECT * FROM youtube_check_rows WHERE "+where+
            " AND youtube_status='BELUM DIPERIKSA' ORDER BY row_no LIMIT ?",params+(CHECK_BATCH,)).fetchall()]
        if not chunk:
            c.close()
            return jsonify(ok=True,done=True,processed=total,total=total)
        # Pemeriksaan via API resmi, satu permintaan untuk hingga 50 video.
        results,error,pause_minutes=check_api_videos([x["video_id"] for x in chunk])
        now=datetime.now(timezone.utc).isoformat()
        if error:
            retry_at=None
            if pause_minutes:
                retry_at=(datetime.now(timezone.utc)+timedelta(minutes=pause_minutes)).isoformat()
                c.execute("""INSERT INTO youtube_check_cooldown(batch_id,retry_at,http_reason,updated_at)
                    VALUES(?,?,?,?) ON CONFLICT(batch_id) DO UPDATE SET
                    retry_at=excluded.retry_at,http_reason=excluded.http_reason,updated_at=excluded.updated_at""",
                    (batch,retry_at,error,now))
            c.commit()
            remaining=c.execute("SELECT COUNT(*) n FROM youtube_check_rows WHERE "+where+
                " AND youtube_status='BELUM DIPERIKSA'",params).fetchone()["n"]
            c.close()
            return jsonify(ok=False,done=False,processed=total-remaining,total=total,
                throttled=bool(retry_at),retry_at=retry_at,error=error,warning=error)
        for x in chunk:
            status,note=results[x["video_id"]]
            c.execute("UPDATE youtube_check_rows SET youtube_status=?,checked_at=?,note=? WHERE batch_id=? AND row_no=?",
                (status,now,note,batch,x["row_no"]))
        c.commit()
        remaining=c.execute("SELECT COUNT(*) n FROM youtube_check_rows WHERE "+where+
            " AND youtube_status='BELUM DIPERIKSA'",params).fetchone()["n"]
        c.close()
        return jsonify(ok=True,done=(remaining==0),processed=total-remaining,total=total,
            throttled=False,retry_at=None,warning="")

    @app.get("/check-video-youtube/export/<batch>")
    @role_required("superadmin","admin_pdn")
    def youtube_check_export(batch):
        user=current_user()
        c=db(); ensure_tables(c); where,params=scoped(batch,user)
        cursor=c.execute("SELECT * FROM youtube_check_rows WHERE "+where+" ORDER BY row_no",params)
        stream=io.StringIO()
        writer=csv.writer(stream)
        writer.writerow(["NO","Judul lagu","artis","status awal","link","video ID","status YouTube","waktu check (UTC)","keterangan"])
        n=0
        for x in cursor:
            n+=1
            writer.writerow([x["source_no"],x["song_title"],x["artist"],x["source_status"],
                x["video_url"],x["video_id"],x["youtube_status"],x["checked_at"],x["note"]])
        c.close()
        if not n: abort(404)
        return Response("\\ufeff"+stream.getvalue(),mimetype="text/csv; charset=utf-8",
            headers={"Content-Disposition":f"attachment; filename=hasil-check-youtube-{batch[:8]}.csv"})

    @app.get("/check-video-youtube/export-unavailable/<batch>")
    @role_required("superadmin","admin_pdn")
    def youtube_check_export_unavailable(batch):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        user=current_user()
        c=db(); ensure_tables(c); where,params=scoped(batch,user)
        exists=c.execute("SELECT COUNT(*) n FROM youtube_check_rows WHERE "+where,params).fetchone()["n"]
        if not exists:
            c.close(); abort(404)
        cursor=c.execute("SELECT source_no,song_title,artist,source_status,video_url,video_id,checked_at,note FROM youtube_check_rows WHERE "+where+
            " AND youtube_status='TIDAK TERSEDIA' ORDER BY row_no",params)
        wb=Workbook(write_only=False)
        ws=wb.active; ws.title="Video Tidak Tersedia"
        headers=["NO","Judul lagu","Artis","Status awal","Link YouTube","Video ID","Status YouTube","Waktu check (UTC)","Keterangan"]
        ws.append(headers)
        for item in cursor:
            ws.append([item["source_no"],item["song_title"],item["artist"],item["source_status"],
                item["video_url"],item["video_id"],"VIDEO TIDAK TERSEDIA",item["checked_at"],item["note"]])
        c.close()
        for cell in ws[1]:
            cell.fill=PatternFill("solid",fgColor="183153")
            cell.font=Font(bold=True,color="FFFFFF")
            cell.alignment=Alignment(wrap_text=True)
        for i,w in enumerate((9,35,25,18,62,19,25,29,65),1):
            ws.column_dimensions[get_column_letter(i)].width=w
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions
        buffer=io.BytesIO();wb.save(buffer);buffer.seek(0)
        from flask import send_file
        return send_file(buffer,as_attachment=True,download_name=f"video-tidak-tersedia-{batch[:8]}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
